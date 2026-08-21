#Import python libraries for specialized functions
import openpyxl
import datetime
import sys

#assign directory string to variable
dir="F:\\Python Code\\Test Scripts\\PythonScriptXL\\Templates\\"
dir2="F:\\Python Code\\Test Scripts\PythonScriptXL\\Reports\\"

#dir='F:\\Python Code\\Test Scripts\\PythonScriptXL\\production\\MEPAG-1428_Example_Excel_Template_Python_Script\\'

#read tempfile data and template workbooks
sourcefile   =openpyxl.load_workbook(dir+"MEPAG_1404_DistributionDEA.xlsx")
templatefile =openpyxl.load_workbook(dir+"MEPAG_1404_Template.xlsx")
template_row=4
template_col=2

from openpyxl import load_workbook
sourcefile_ws = load_workbook(filename=dir+"MEPAG_1404_DistributionDEA.xlsx")
sourcefile_ws.sheetnames [0]
sourceworksheet = sourcefile_ws.sheetnames [0]
print('sheet name from data file is:' + sourcefile_ws.sheetnames [0])

templatefile_ws = load_workbook(filename=dir+"MEPAG_1404_Template.xlsx")
templatefile_ws.sheetnames [1]
templatesheet = templatefile_ws.sheetnames [1]
print('1st sheet name from template is:' + templatefile_ws.sheetnames [0])
print('2nd sheet name from template is:' + templatefile_ws.sheetnames [1])
template_row=4
template_col=0


#loop through worksheets of tempfile
for sheetname in sourcefile.sheetnames:
    sourcefile_ws=sourcefile[sheetname]
    templatefile_ws=templatefile[sheetname]

    #loop through columns of tempfile worksheet
    for colnum in range(1, sourcefile_ws.max_column + 1):
        print(sourcefile_ws.cell(1, column = colnum).value)
        if sourcefile_ws.cell(1, column = colnum).value == 'processdate':
            print('Process Date: ' + str(sourcefile_ws.cell(2, column = colnum).value))
            templatefile_ws.cell(1, 1).value = 'Process Date: ' + str(sourcefile_ws.cell(2, column = colnum).value)
            templatefile_ws.cell(2, 1).value = 'Blue Cross Blue Shield of Michigan (CustomerID 13) ' + str(sourcefile_ws.cell(2, 2).value)
            templatefile_ws.cell(3, 1).value = 'Contract Profiles (238, 393, 8634) and HNUM (H9572, S5584)'

#loop through worksheets of tempfile
for sheetname2 in sourcefile.sheetnames:
    sourcefile_ws2=sourcefile[sheetname2]
    templatefile_ws2=templatefile[sheetname2]

    template_row=4


    for colnum2 in range(2, sourcefile_ws2.max_column + 1):
        template_row=4
        template_col=template_col+1
        #loop through rows of tempfile worksheet starting at second row (exclude header row)
        for rownum2 in range(2, sourcefile_ws2.max_row + 1):
            template_row=template_row+1
            templatefile_ws2.cell(row = template_row, column = template_col).value = sourcefile_ws2.cell(row = rownum2, column = colnum2).value
            print(sourcefile_ws2.cell(row = rownum2, column = colnum2).value)
            
            #templatefile_ws.cell(row = template_row+1, column = colnum).value = sourcefile_ws.cell(row =rownum, column = colnum)

#assign today's date
nowdate=datetime.datetime.now().strftime("%d%b%Y").upper()

#save report with optional pass-thru variable(s) or today's date
# template.save(dir+'Reports\\Year_'+rpc_year+'_RPCs_and_Approvals_'+nowdate+'.xlsx')

#save report with optional pass-thru variable(s) or today's date
templatefile.save(dir2+'MEPAG_1404_DistributionDEA_Report'+nowdate+'.xlsx')
