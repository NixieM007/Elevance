import pandas as pd
import numpy as np
import pandas.io.sql
import pyodbc
import xlrd
import openpyxl
import datetime
import sys

r=2

server = "VMSQLDWBIDEV"
db = 'Sandbox_ConceptDevelopment'
conn = pyodbc.connect('DRIVER={SQL Server};SERVER=' + server + ';DATABASE=' + 
db + ';Trusted_Connection=yes')


cursor = conn.cursor()
#read tempfile data and template workbooks

#book = xlrd.open_workbook_xls("F:\Vendors\Pareto\MA ESRD and MSP Premium Recovery Validation\ValidatedFiles\Pareto_MonthlyRecoveries_Validated_August 2023.xlsx") 
#sheet = book.sheet_by_name("MSP MMR Recoveries EDW")


#from openpyxl import load_workbook

sourcefile   = openpyxl.load_workbook(filename="F:\Vendors\Pareto\MA ESRD and MSP Premium Recovery Validation\ValidatedFiles\Pareto_MonthlyRecoveries_Validated_August 2023.xlsx")


query = """
INSERT INTO [dbo].[NM_PI_VendorRecoveries_ParetoMSPESRD]
           ([HICN_MCOContract]
           ,[HICNumber]
           ,[PaymentDate]
           ,[PaymtAdjustmentMSAStartDate]
           ,[PaymtAdjustmentMSAEndDate]
           ,[AdjustmentReasonCode]
           ,[RecordType]
           ,[TotalPartCPayment]
           ,[NumberofPaymtAdjustmtMonthsPartA]
           ,[NumberofPaymtAdjustmtMonthsPartB]
           ,[MCOContractNumber]) 
           VALUES  (?, ?, ?, ?, ?, ?, ?, 
                    ?, ?, ?,?)  """

#loop through worksheets of tempfile
for sheetname in sourcefile.sheetnames:
    sourceworksheet = sheetname
    if sheetname == 'MSP MMR Recoveries EDW' or sheetname == 'ESRD MMR Recoveries EDW':
        # print('sheet name from data file is:' + sourcefile_ws.sheetnames [1])
        print('sheet name from data file is:' + sheetname)

        for r in  range(2, sourcefile[sheetname].max_row):
                HICN_MCOContract                    = sourcefile[sheetname].cell(r,1).value
                HICNumber                           = sourcefile[sheetname].cell(r,1).value
                PaymentDate                         = sourcefile[sheetname].cell(r,2).value
                PaymtAdjustmentMSAStartDate         = sourcefile[sheetname].cell(r,3).value
                PaymtAdjustmentMSAEndDate           = sourcefile[sheetname].cell(r,4).value
                AdjustmentReasonCode                = sourcefile[sheetname].cell(r,5).value
                RecordType                          = sourcefile[sheetname].cell(r,6).value
                TotalPartCPayment                   = sourcefile[sheetname].cell(r,7).value
                NumberofPaymtAdjustmtMonthsPartA    = sourcefile[sheetname].cell(r,8).value
                NumberofPaymtAdjustmtMonthsPartB    = sourcefile[sheetname].cell(r,9).value
                MCOContractNumber                   = sourcefile[sheetname].cell(r,10).value
                
                
                print(sourcefile[sheetname].cell(r,1).value, sourcefile[sheetname].cell(r,1).value, sourcefile[sheetname].cell(r,2).value,
                sourcefile[sheetname].cell(r,3).value, sourcefile[sheetname].cell(r,4).value, sourcefile[sheetname].cell(r,5).value, sourcefile[sheetname].cell(r,6).value
                , sourcefile[sheetname].cell(r,7).value, sourcefile[sheetname].cell(r,8).value, sourcefile[sheetname].cell(r,9).value, sourcefile[sheetname].cell(r,10).value, sourcefile[sheetname].cell(r,6).value)
                
                
                print(HICN_MCOContract, HICNumber,
                PaymentDate ,
                PaymtAdjustmentMSAStartDate,
                PaymtAdjustmentMSAEndDate,
                AdjustmentReasonCode,
                RecordType,
                TotalPartCPayment,
                NumberofPaymtAdjustmtMonthsPartA,
                NumberofPaymtAdjustmtMonthsPartB,
                MCOContractNumber)
        
                values = (HICN_MCOContract, HICNumber, PaymentDate, PaymtAdjustmentMSAStartDate, 
                            PaymtAdjustmentMSAEndDate, AdjustmentReasonCode, RecordType, TotalPartCPayment, NumberofPaymtAdjustmtMonthsPartA, NumberofPaymtAdjustmtMonthsPartB, 
                            MCOContractNumber)
                
                #cursor.execute(query, values)
                #db.commit()
                #db.close()