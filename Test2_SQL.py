import pandas as pd
import numpy as np
import pandas.io.sql
import pyodbc
import xlrd
import openpyxl
import datetime
import sys
import string
import arrow

# BOTH ESRD AND MSP MUST BE RUN IN THE SAME SCRIPT. COPY AND PASTE BOTH MEDICARE NUMBERS IN BEFORE RUN
# CHANGE FILENAME MONTH NUMBER

strValidationDate = datetime.datetime.now().strftime("%d%b%Y").upper()

#from openpyxl import load_workbook
strfilename     = "F:\Vendors\Pareto\MA ESRD and MSP Premium Recovery Validation\ValidatedFiles\Pareto_MonthlyRecoveries_Validated_082023.xlsx"
sourcefile      = openpyxl.load_workbook(filename=strfilename)
strfilemmyyyy   = strfilename[112:-5]
strfilemm       = strfilename[112:-9]
strfileyyyy     = strfilename[114:-5]
strfiledate     = strfileyyyy + strfilemm + '01'
strnowdated     = arrow.now().format('YYYYMMDD')
Recondate       = strnowdated
strVendorName   = "Pareto Intelligence" 
strCMSMBI      = "('5WC6GV1VH65','5HF1M88HU28','7QE1P61JW38')"
#strCMSMBI      = "('9JA5MR9YF63')"
strLOB          = "Medicare Advantage (MA)"
r = 2
sourcefile.create_sheet('MMR Recoveries EDW')
sourcesheet = sourcefile['MMR Recoveries EDW']
#print(sourcesheet, sourcefile)

sourcesheet.cell(1,1).value    =  "HICN_MCOContract"
sourcesheet.cell(1,2).value    =  "HICNumber"
sourcesheet.cell(1,3).value    =  "PaymentDate"
sourcesheet.cell(1,4).value    =  "PaymtAdjustmentMSAStartDate"
sourcesheet.cell(1,5).value    =  "PaymtAdjustmentMSAEndDate"
sourcesheet.cell(1,6).value    =  "AdjustmentReasonCode"
sourcesheet.cell(1,7).value    =  "RecordType"
sourcesheet.cell(1,8).value    =  "TotalPartCPayment"               
sourcesheet.cell(1,9).value    =  "NumberofPaymtAdjustmtMonthsPartA"
sourcesheet.cell(1,10).value   =  "NumberofPaymtAdjustmtMonthsPartB"
sourcesheet.cell(1,11).value   =  "MCOContractNumber"



serverprod = "JHHCSQLDWBI"
dbprod = 'BI_DataLake_PROD'
connprod = pyodbc.connect('DRIVER={SQL Server};SERVER=' + serverprod + ';DATABASE=' + dbprod + ';Trusted_Connection=yes')

cursorprod = connprod.cursor()

queryprod = """ select HICNumber+MCOContractNumber, HICNumber, PaymentDate, min(PaymtAdjustmentMSAStartDate),max(PaymtAdjustmentMSAEndDate),
                AdjustmentReasonCode, case when ESRD = 'Y' then 'ESRD' else 'MSP' end, sum(TotalPartCPayment),sum(NumberofPaymtAdjustmtMonthsPartA), 
                sum(NumberofPaymtAdjustmtMonthsPartB), MCOContractNumber FROM [BI_DataLake_PROD].MA.CMMS_MonthlyMembership_Detail 
                WHERE HICNumber in """ +  strCMSMBI + """ and AdjustmentReasonCode in ('08', '42') and PaymentDate >= '202211' 
                group by HICNumber, PaymentDate, AdjustmentReasonCode, ESRD, MCOContractNumber  """ 

row = cursorprod.execute (queryprod)
rows = cursorprod.fetchall()
for row in rows:
        print (row)
        
             
        serverdev = 'VMSQLDWBIDEV'
        dbdev = 'Sandbox_ConceptDevelopment'
        conndev = pyodbc.connect('DRIVER={SQL Server};SERVER=' + serverdev + ';DATABASE=' + 
        dbdev + ';Trusted_Connection=yes')
        
        VendorName                              = strVendorName
        LineofBusinessName                      = strLOB 
        strValidationDate                       = strnowdated  
        VendorFileDate                          = strfiledate
        HICN_MCOContract                        = row[0]
        HICNumber                               = row[1]
        PaymentDate                             = row[2]
        PaymtAdjustmentMSAStartDate             = row[3]
        PaymtAdjustmentMSAEndDate               = row[4]
        AdjustmentReasonCode                    = row[5]
        RecordType                              = row[6]
        TotalPartCPayment                       = row[7]
        NumberofPaymtAdjustmtMonthsPartA        = row[8]
        NumberofPaymtAdjustmtMonthsPartB        = row[9]
        MCOContractNumber                       = row[10]
                       
  
        
        querydev = """
                INSERT INTO [dbo].[NM_PI_VendorRecoveries_ParetoMSPESRD]
                (VendorName, LineofBusinessName, ValidationDate, VendorFileDate, HICN_MCOContract, HICNumber, PaymentDate, PaymtAdjustmentMSAStartDate, 
                PaymtAdjustmentMSAEndDate,AdjustmentReasonCode, RecordType, TotalPartCPayment, NumberofPaymtAdjustmtMonthsPartA, NumberofPaymtAdjustmtMonthsPartB, MCOContractNumber) 
                VALUES  (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)  """
        
        values = (VendorName, LineofBusinessName, strValidationDate, VendorFileDate, HICN_MCOContract, HICNumber, PaymentDate, PaymtAdjustmentMSAStartDate, 
                  PaymtAdjustmentMSAEndDate, AdjustmentReasonCode, RecordType, TotalPartCPayment, NumberofPaymtAdjustmtMonthsPartA, NumberofPaymtAdjustmtMonthsPartB, MCOContractNumber)
        
        
        sourcesheet.cell(r,1).value  =  HICN_MCOContract
        sourcesheet.cell(r,2).value  =  HICNumber
        sourcesheet.cell(r,3).value  =  PaymentDate
        sourcesheet.cell(r,4).value  =  PaymtAdjustmentMSAStartDate
        sourcesheet.cell(r,5).value  =  PaymtAdjustmentMSAEndDate
        sourcesheet.cell(r,6).value  =  AdjustmentReasonCode
        sourcesheet.cell(r,7).value  =  RecordType
        sourcesheet.cell(r,8).value  =  TotalPartCPayment
        sourcesheet.cell(r,9).value  =  NumberofPaymtAdjustmtMonthsPartA
        sourcesheet.cell(r,10).value =  NumberofPaymtAdjustmtMonthsPartB
        sourcesheet.cell(r,11).value =  MCOContractNumber
        
        r = r + 1
        
        cursordev = conndev.cursor()
        cursordev.execute(querydev, values)
        cursordev.commit()
        cursordev.close()
        
        
#Save Workbook
sourcefile.save(strfilename)

