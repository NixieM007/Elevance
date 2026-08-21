import pandas as pd
import numpy as np
import pandas.io.sql
import pyodbc
import xlrd
import openpyxl
import datetime
import sys
import string
import arrow

# BOTH ESRD AND MSP MUST BE RUN IN THE SAME SCRIPT. COPY AND PASTE BOTH MEDICARE NUMBERS IN BEFORE RUN
# CHANGE FILENAME MONTH NUMBER

strValidationDate = datetime.datetime.now().strftime("%d%b%Y").upper()

#from openpyxl import load_workbook
strfilename     = "F:\Vendors\Pareto\MA ESRD and MSP Premium Recovery Validation\ValidatedFiles\Pareto_MonthlyRecoveries_Validated_082023.xlsx"
sourcefile      = openpyxl.load_workbook(filename=strfilename)
strfilemmyyyy   = strfilename[112:-5]
strfilemm       = strfilename[112:-9]
strfileyyyy     = strfilename[114:-5]
strfiledate     = strfileyyyy + strfilemm + '01'
strnowdated     = arrow.now().format('YYYYMMDD')
Recondate       = strnowdated
strVendorName   = "Pareto Intelligence" 
strCMSMBI      = "('1WR2DW6MV98','2TA9MY4NU24','3GP9E27QD32','7J57DH6YM53','9C85DT0CM96')"
#strCMSMBI      = "('9JA5MR9YF63')"
strLOB          = "Medicare Advantage (MA)"
r = 2
sourcefile.create_sheet('MMR Recoveries EDW')
sourcesheet = sourcefile['MMR Recoveries EDW']
#print(sourcesheet, sourcefile)





serverprod = "JHHCSQLDWBI"
dbprod = 'BI_DataLake_PROD'
connprod = pyodbc.connect('DRIVER={SQL Server};SERVER=' + serverprod + ';DATABASE=' + dbprod + ';Trusted_Connection=yes')

cursorprod = connprod.cursor()

queryprod = """ select HICNumber+MCOContractNumber, HICNumber, PaymentDate, min(PaymtAdjustmentMSAStartDate),max(PaymtAdjustmentMSAEndDate),
                AdjustmentReasonCode, case when ESRD = 'Y' then 'ESRD' else 'MSP' end, sum(TotalPartCPayment),sum(NumberofPaymtAdjustmtMonthsPartA), 
                sum(NumberofPaymtAdjustmtMonthsPartB), MCOContractNumber FROM [BI_DataLake_PROD].MA.CMMS_MonthlyMembership_Detail 
                WHERE HICNumber in """ +  strCMSMBI + """ and AdjustmentReasonCode in ('08', '42') and PaymentDate >= '202211' 
                group by HICNumber, PaymentDate, AdjustmentReasonCode, ESRD, MCOContractNumber  """ 

row = cursorprod.execute (queryprod)
rows = cursorprod.fetchall()
for row in rows:
        print (row)
        
             
        serverdev = 'VMSQLDWBIDEV'
        dbdev = 'Sandbox_ConceptDevelopment'
        conndev = pyodbc.connect('DRIVER={SQL Server};SERVER=' + serverdev + ';DATABASE=' + 
        dbdev + ';Trusted_Connection=yes')
        
        VendorName                              = strVendorName
        LineofBusinessName                      = strLOB 
        strValidationDate                       = strnowdated  
        VendorFileDate                          = strfiledate
        HICN_MCOContract                        = row[0]
        HICNumber                               = row[1]
        PaymentDate                             = row[2]
        PaymtAdjustmentMSAStartDate             = row[3]
        PaymtAdjustmentMSAEndDate               = row[4]
        AdjustmentReasonCode                    = row[5]
        RecordType                              = row[6]
        TotalPartCPayment                       = row[7]
        NumberofPaymtAdjustmtMonthsPartA        = row[8]
        NumberofPaymtAdjustmtMonthsPartB        = row[9]
        MCOContractNumber                       = row[10]
                       
  
        
        querydev = """
                INSERT INTO [dbo].[NM_PI_VendorRecoveries_ParetoMSPESRD]
                (VendorName, LineofBusinessName, ValidationDate, VendorFileDate, HICN_MCOContract, HICNumber, PaymentDate, PaymtAdjustmentMSAStartDate, 
                PaymtAdjustmentMSAEndDate,AdjustmentReasonCode, RecordType, TotalPartCPayment, NumberofPaymtAdjustmtMonthsPartA, NumberofPaymtAdjustmtMonthsPartB, MCOContractNumber) 
                VALUES  (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)  """
        
        values = (VendorName, LineofBusinessName, strValidationDate, VendorFileDate, HICN_MCOContract, HICNumber, PaymentDate, PaymtAdjustmentMSAStartDate, 
                  PaymtAdjustmentMSAEndDate, AdjustmentReasonCode, RecordType, TotalPartCPayment, NumberofPaymtAdjustmtMonthsPartA, NumberofPaymtAdjustmtMonthsPartB, MCOContractNumber)
        
        
        
        
        cursordev = conndev.cursor()
        cursordev.execute(querydev, values)
        cursordev.commit()
        cursordev.close()
        

